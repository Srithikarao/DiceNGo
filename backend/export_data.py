import os
import json
import openpyxl

wb = openpyxl.load_workbook('backend/data/warangal_database.xlsx', data_only=True)

# 1. Food
food = []
if "Food" in wb.sheetnames:
    ws = wb["Food"]
    rows = list(ws.iter_rows(values_only=True))
    for idx, r in enumerate(rows[1:], start=1):
        if not r or not r[1]: continue
        food.append({
            "id": idx,
            "name": str(r[1]).strip(),
            "category": str(r[2]).strip(),
            "cuisine": str(r[3] or ""),
            "address": str(r[4] or ""),
            "area": str(r[5] or "Hanamkonda"),
            "latitude": float(r[6] or 18.0050),
            "longitude": float(r[7] or 79.5600),
            "maps_url": str(r[8] or ""),
            "opening_time": str(r[9] or "11:00 AM"),
            "closing_time": str(r[10] or "11:00 PM"),
            "rating": float(r[11] or 4.0),
            "review_count": int(r[12] or 0),
            "price_range": str(r[13] or "₹₹"),
            "best_known_for": str(r[14] or "Specialty"),
            "veg": (str(r[15]).upper() == "YES"),
            "non_veg": (str(r[16]).upper() == "YES"),
            "indoor_seating": (str(r[17]).upper() == "YES"),
            "outdoor_seating": (str(r[18]).upper() == "YES"),
            "parking": (str(r[19]).upper() == "YES"),
            "discounts": str(r[20] or "None"),
            "seating_capacity": int(r[21] or 40),
            "instagram_url": str(r[22] or ""),
            "images": str(r[23] or ""),
            "description": str(r[24] or ""),
            "is_new": (hash(str(r[1])) % 7 == 0),
            "is_featured": (hash(str(r[1])) % 9 == 0)
        })

# 2. Explore
explore = []
if "Explore" in wb.sheetnames:
    ws = wb["Explore"]
    rows = list(ws.iter_rows(values_only=True))
    for idx, r in enumerate(rows[1:], start=1):
        if not r or not r[1]: continue
        explore.append({
            "id": idx,
            "name": str(r[1]).strip(),
            "category": str(r[2]).strip(),
            "address": str(r[3] or ""),
            "area": str(r[4] or "Warangal"),
            "latitude": float(r[5] or 18.0019),
            "longitude": float(r[6] or 79.5781),
            "maps_url": str(r[7] or ""),
            "best_time": str(r[8] or "Evening"),
            "entry_fee": str(r[9] or "Free"),
            "opening_time": str(r[10] or "06:00 AM"),
            "closing_time": str(r[11] or "08:00 PM"),
            "rating": float(r[12] or 4.2),
            "review_count": int(r[13] or 0),
            "parking": (str(r[14]).upper() == "YES"),
            "friends_suitable": (str(r[15]).upper() == "YES"),
            "family_suitable": (str(r[16]).upper() == "YES"),
            "images": str(r[17] or ""),
            "description": str(r[18] or ""),
            "is_new": (hash(str(r[1])) % 6 == 0),
            "is_featured": (hash(str(r[1])) % 8 == 0)
        })

# 3. Events
events = []
if "Events" in wb.sheetnames:
    ws = wb["Events"]
    rows = list(ws.iter_rows(values_only=True))
    for idx, r in enumerate(rows[1:], start=1):
        if not r or not r[1]: continue
        events.append({
            "id": idx,
            "event_name": str(r[1]).strip(),
            "venue": str(r[2] or ""),
            "category": str(r[3] or "Other"),
            "address": str(r[4] or ""),
            "area": str(r[5] or "Warangal"),
            "latitude": float(r[6] or 18.0045),
            "longitude": float(r[7] or 79.5391),
            "date": str(r[8] or "2026-09-20"),
            "start_time": str(r[9] or "05:00 PM"),
            "end_time": str(r[10] or "09:00 PM"),
            "individual_or_group": str(r[11] or "Group"),
            "contact": str(r[12] or ""),
            "maps_url": str(r[13] or ""),
            "instagram_url": str(r[14] or ""),
            "website_url": str(r[15] or ""),
            "fee": str(r[16] or "Free"),
            "rating": float(r[17] or 4.5),
            "reviews": int(r[18] or 0),
            "images": str(r[19] or ""),
            "description": str(r[20] or ""),
            "booking_information": str(r[21] or "Contact venue"),
            "is_featured": True
        })

missions = [
    {"id": 1, "title": "Biryani Quest", "description": "Visit 1 Biryani spot or authentic mess in Hanamkonda today.", "mission_type": "Food", "target_category": "Biryani", "reward_points": 100},
    {"id": 2, "title": "Sunset Chaser", "description": "Discover a scenic lake bund or hilltop viewpoint before 6:45 PM.", "mission_type": "Explore", "target_category": "Sunset Spots", "reward_points": 150},
    {"id": 3, "title": "Chai & Banter Adda", "description": "Hangout at a student cafe or Irani tea adda with your gang.", "mission_type": "Food", "target_category": "Cafes", "reward_points": 75},
    {"id": 4, "title": "Heritage Unlock", "description": "Explore a Kakatiya monument or heritage temple you haven't visited.", "mission_type": "Explore", "target_category": "Temples", "reward_points": 200},
]

os.makedirs('frontend/src/data', exist_ok=True)
output_path = 'frontend/src/data/warangalData.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump({
        "food": food,
        "explore": explore,
        "events": events,
        "missions": missions
    }, f, ensure_ascii=False, indent=2)

print(f"Exported {len(food)} food places, {len(explore)} explore places, {len(events)} events to {output_path}")
