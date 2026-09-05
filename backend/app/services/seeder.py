import os
import openpyxl
from sqlalchemy.orm import Session
from backend.app.models import FoodPlace, ExplorePlace, EventItem, DailyMission

# Path to warangal_database.xlsx: backend/data/warangal_database.xlsx
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__)) # backend/app/services
APP_DIR = os.path.dirname(CURRENT_DIR) # backend/app
BACKEND_DIR = os.path.dirname(APP_DIR) # backend
EXCEL_PATH = os.path.join(BACKEND_DIR, "data", "warangal_database.xlsx")

def seed_database_from_excel(db: Session, force_reload: bool = False):
    if not os.path.exists(EXCEL_PATH):
        print(f"[SEEDER WARNING] Excel file not found at {EXCEL_PATH}")
        return

    food_count = db.query(FoodPlace).count()
    explore_count = db.query(ExplorePlace).count()
    event_count = db.query(EventItem).count()

    if not force_reload and food_count >= 300 and explore_count >= 100 and event_count >= 10:
        print(f"[SEEDER] Database already initialized ({food_count} food, {explore_count} explore, {event_count} events).")
        seed_missions_if_needed(db)
        return

    print(f"[SEEDER] Loading workbook from {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # 1. SEED FOOD
    if force_reload or food_count < 300:
        if force_reload:
            db.query(FoodPlace).delete()
        if "Food" in wb.sheetnames:
            ws = wb["Food"]
            rows = list(ws.iter_rows(values_only=True))
            header = rows[0]
            print(f"[SEEDER] Seeding {len(rows)-1} Food records...")
            for r in rows[1:]:
                if not r or not r[1]: continue
                food = FoodPlace(
                    name=str(r[1]).strip(),
                    category=str(r[2]).strip(),
                    cuisine=str(r[3] or ""),
                    address=str(r[4] or ""),
                    area=str(r[5] or "Hanamkonda"),
                    latitude=float(r[6] or 18.0050),
                    longitude=float(r[7] or 79.5600),
                    maps_url=str(r[8] or ""),
                    opening_time=str(r[9] or "11:00 AM"),
                    closing_time=str(r[10] or "11:00 PM"),
                    rating=float(r[11] or 4.0),
                    review_count=int(r[12] or 0),
                    price_range=str(r[13] or "₹₹"),
                    best_known_for=str(r[14] or "Specialty"),
                    veg=(str(r[15]).upper() == "YES"),
                    non_veg=(str(r[16]).upper() == "YES"),
                    indoor_seating=(str(r[17]).upper() == "YES"),
                    outdoor_seating=(str(r[18]).upper() == "YES"),
                    parking=(str(r[19]).upper() == "YES"),
                    discounts=str(r[20] or "None"),
                    seating_capacity=int(r[21] or 40),
                    instagram_url=str(r[22] or ""),
                    images=str(r[23] or ""),
                    description=str(r[24] or ""),
                    is_new=(hash(str(r[1])) % 7 == 0),
                    is_featured=(hash(str(r[1])) % 9 == 0)
                )
                db.add(food)
            db.commit()

    # 2. SEED EXPLORE
    if force_reload or explore_count < 100:
        if force_reload:
            db.query(ExplorePlace).delete()
        if "Explore" in wb.sheetnames:
            ws = wb["Explore"]
            rows = list(ws.iter_rows(values_only=True))
            print(f"[SEEDER] Seeding {len(rows)-1} Explore records...")
            for r in rows[1:]:
                if not r or not r[1]: continue
                explore = ExplorePlace(
                    name=str(r[1]).strip(),
                    category=str(r[2]).strip(),
                    address=str(r[3] or ""),
                    area=str(r[4] or "Warangal"),
                    latitude=float(r[5] or 18.0019),
                    longitude=float(r[6] or 79.5781),
                    maps_url=str(r[7] or ""),
                    best_time=str(r[8] or "Evening"),
                    entry_fee=str(r[9] or "Free"),
                    opening_time=str(r[10] or "06:00 AM"),
                    closing_time=str(r[11] or "08:00 PM"),
                    rating=float(r[12] or 4.2),
                    review_count=int(r[13] or 0),
                    parking=(str(r[14]).upper() == "YES"),
                    friends_suitable=(str(r[15]).upper() == "YES"),
                    family_suitable=(str(r[16]).upper() == "YES"),
                    images=str(r[17] or ""),
                    description=str(r[18] or ""),
                    is_new=(hash(str(r[1])) % 6 == 0),
                    is_featured=(hash(str(r[1])) % 8 == 0)
                )
                db.add(explore)
            db.commit()

    # 3. SEED EVENTS
    if force_reload or event_count < 10:
        if force_reload:
            db.query(EventItem).delete()
        if "Events" in wb.sheetnames:
            ws = wb["Events"]
            rows = list(ws.iter_rows(values_only=True))
            print(f"[SEEDER] Seeding {len(rows)-1} Events...")
            for r in rows[1:]:
                if not r or not r[1]: continue
                event = EventItem(
                    event_name=str(r[1]).strip(),
                    venue=str(r[2] or ""),
                    category=str(r[3] or "Other"),
                    address=str(r[4] or ""),
                    area=str(r[5] or "Warangal"),
                    latitude=float(r[6] or 18.0045),
                    longitude=float(r[7] or 79.5391),
                    date=str(r[8] or "2026-09-20"),
                    start_time=str(r[9] or "05:00 PM"),
                    end_time=str(r[10] or "09:00 PM"),
                    individual_or_group=str(r[11] or "Group"),
                    contact=str(r[12] or ""),
                    maps_url=str(r[13] or ""),
                    instagram_url=str(r[14] or ""),
                    website_url=str(r[15] or ""),
                    fee=str(r[16] or "Free"),
                    rating=float(r[17] or 4.5),
                    reviews=int(r[18] or 0),
                    images=str(r[19] or ""),
                    description=str(r[20] or ""),
                    booking_information=str(r[21] or "Contact venue"),
                    is_featured=True
                )
                db.add(event)
            db.commit()

    seed_missions_if_needed(db)
    print(f"[SEEDER COMPLETE] {db.query(FoodPlace).count()} Food, {db.query(ExplorePlace).count()} Explore, {db.query(EventItem).count()} Events loaded.")

def seed_missions_if_needed(db: Session):
    if db.query(DailyMission).count() == 0:
        missions = [
            DailyMission(title="Biryani Quest", description="Visit 1 Biryani spot or authentic mess in Hanamkonda today.", mission_type="Food", target_category="Biryani", reward_points=100),
            DailyMission(title="Sunset Chaser", description="Discover a scenic lake bund or hilltop viewpoint before 6:45 PM.", mission_type="Explore", target_category="Sunset Spots", reward_points=150),
            DailyMission(title="Chai & Banter Adda", description="Hangout at a student cafe or Irani tea adda with your gang.", mission_type="Food", target_category="Cafes", reward_points=75),
            DailyMission(title="Heritage Unlock", description="Explore a Kakatiya monument or heritage temple you haven't visited.", mission_type="Explore", target_category="Temples", reward_points=200),
        ]
        for m in missions:
            db.add(m)
        db.commit()
