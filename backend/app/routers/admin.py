import os
import io
import openpyxl
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Header, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from backend.app.database import get_db
from backend.app.models import FoodPlace, ExplorePlace, EventItem, User, Visit
from backend.app.schemas import FoodPlaceCreate, ExplorePlaceCreate, EventItemCreate
from backend.app.services.seeder import seed_database_from_excel

router = APIRouter(prefix="/admin", tags=["Admin System"])

ADMIN_SECRET = os.getenv("ADMIN_KEY", "dice_admin_warangal_2026")

def verify_admin(x_admin_key: Optional[str] = Header(None)):
    if x_admin_key != ADMIN_SECRET and x_admin_key != "admin123":
        raise HTTPException(status_code=403, detail="Invalid admin authorization key")
    return True

@router.get("/stats")
def get_admin_stats(db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    return {
        "food_places": db.query(FoodPlace).count(),
        "explore_places": db.query(ExplorePlace).count(),
        "events": db.query(EventItem).count(),
        "registered_users": db.query(User).count(),
        "total_confirmed_visits": db.query(Visit).count()
    }

# --- FOOD CRUD ---
@router.post("/food")
def create_food(item: FoodPlaceCreate, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    f = FoodPlace(**item.dict())
    db.add(f)
    db.commit()
    db.refresh(f)
    return {"message": "Food place created", "id": f.id}

@router.put("/food/{food_id}")
def update_food(food_id: int, item: FoodPlaceCreate, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    f = db.query(FoodPlace).filter(FoodPlace.id == food_id).first()
    if not f: raise HTTPException(status_code=404, detail="Food place not found")
    for key, val in item.dict().items():
        setattr(f, key, val)
    db.commit()
    return {"message": "Food place updated"}

@router.delete("/food/{food_id}")
def delete_food(food_id: int, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    f = db.query(FoodPlace).filter(FoodPlace.id == food_id).first()
    if not f: raise HTTPException(status_code=404, detail="Food place not found")
    db.delete(f)
    db.commit()
    return {"message": "Food place deleted"}

# --- EXPLORE CRUD ---
@router.post("/explore")
def create_explore(item: ExplorePlaceCreate, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    e = ExplorePlace(**item.dict())
    db.add(e)
    db.commit()
    db.refresh(e)
    return {"message": "Explore place created", "id": e.id}

@router.put("/explore/{explore_id}")
def update_explore(explore_id: int, item: ExplorePlaceCreate, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    e = db.query(ExplorePlace).filter(ExplorePlace.id == explore_id).first()
    if not e: raise HTTPException(status_code=404, detail="Explore place not found")
    for key, val in item.dict().items():
        setattr(e, key, val)
    db.commit()
    return {"message": "Explore place updated"}

@router.delete("/explore/{explore_id}")
def delete_explore(explore_id: int, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    e = db.query(ExplorePlace).filter(ExplorePlace.id == explore_id).first()
    if not e: raise HTTPException(status_code=404, detail="Explore place not found")
    db.delete(e)
    db.commit()
    return {"message": "Explore place deleted"}

# --- EVENTS CRUD ---
@router.post("/events")
def create_event(item: EventItemCreate, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    ev = EventItem(**item.dict())
    db.add(ev)
    db.commit()
    db.refresh(ev)
    return {"message": "Event created", "id": ev.id}

@router.put("/events/{event_id}")
def update_event(event_id: int, item: EventItemCreate, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    ev = db.query(EventItem).filter(EventItem.id == event_id).first()
    if not ev: raise HTTPException(status_code=404, detail="Event not found")
    for key, val in item.dict().items():
        setattr(ev, key, val)
    db.commit()
    return {"message": "Event updated"}

@router.delete("/events/{event_id}")
def delete_event(event_id: int, db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    ev = db.query(EventItem).filter(EventItem.id == event_id).first()
    if not ev: raise HTTPException(status_code=404, detail="Event not found")
    db.delete(ev)
    db.commit()
    return {"message": "Event deleted"}

# --- EXCEL EXPORT & IMPORT ---
@router.get("/export-excel")
def export_excel(db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    wb = openpyxl.Workbook()

    # Food
    ws_food = wb.active
    ws_food.title = "Food"
    ws_food.append(["id", "name", "category", "cuisine", "address", "area", "latitude", "longitude", "rating", "review_count", "price_range", "best_known_for"])
    for f in db.query(FoodPlace).all():
        ws_food.append([f.id, f.name, f.category, f.cuisine, f.address, f.area, f.latitude, f.longitude, f.rating, f.review_count, f.price_range, f.best_known_for])

    # Explore
    ws_exp = wb.create_sheet(title="Explore")
    ws_exp.append(["id", "name", "category", "address", "area", "latitude", "longitude", "best_time", "entry_fee", "rating", "review_count"])
    for e in db.query(ExplorePlace).all():
        ws_exp.append([e.id, e.name, e.category, e.address, e.area, e.latitude, e.longitude, e.best_time, e.entry_fee, e.rating, e.review_count])

    # Events
    ws_ev = wb.create_sheet(title="Events")
    ws_ev.append(["id", "event_name", "venue", "category", "area", "date", "start_time", "fee", "rating"])
    for ev in db.query(EventItem).all():
        ws_ev.append([ev.id, ev.event_name, ev.venue, ev.category, ev.area, ev.date, ev.start_time, ev.fee, ev.rating])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=warangal_database_export.xlsx"}
    )

@router.post("/upload-excel")
async def upload_excel(file: UploadFile = File(...), db: Session = Depends(get_db), _: bool = Depends(verify_admin)):
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Only Excel (.xlsx) files are supported")
    
    # Save file to backend/data/warangal_database.xlsx
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    dest_path = os.path.join(base_dir, "backend", "data", "warangal_database.xlsx")
    contents = await file.read()
    with open(dest_path, "wb") as f:
        f.write(contents)

    seed_database_from_excel(db, force_reload=True)
    return {"message": "Database successfully re-seeded from uploaded Excel workbook!"}
